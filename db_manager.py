# db_manager.py
import sqlite3
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, Any, List

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from config import IMPRESORAS, CONFIG, NOMBRES, SECTORES

logger = logging.getLogger(__name__)

DB_FILE = "contadores.db"

def init_db():
    """Crea la tabla registros si no existe."""
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS registros (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ip TEXT NOT NULL,
            nombre TEXT NOT NULL,
            sector TEXT NOT NULL,
            fecha TEXT NOT NULL,  -- YYYY-MM-DD
            contador INTEGER,
            diferencia INTEGER,
            estado TEXT
        )
    """)
    conn.commit()
    conn.close()
    logger.info("Base de datos SQLite lista.")

def tiene_datos():
    """Verifica si la base de datos ya tiene registros."""
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM registros")
    count = c.fetchone()[0]
    conn.close()
    return count > 0

def guardar_registro(ip: str, nombre: str, sector: str, fecha: str, 
                     contador: Optional[int], diferencia: Optional[int]) -> None:
    """Inserta un registro en SQLite."""
    estado = "OK" if contador is not None else "Sin conexión"
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("""
        INSERT INTO registros (ip, nombre, sector, fecha, contador, diferencia, estado)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (ip, nombre, sector, fecha, contador, diferencia, estado))
    conn.commit()
    conn.close()
    logger.info(f"Guardado: {ip} fecha {fecha} contador {contador} dif {diferencia}")

def obtener_ultimo_registro(ip: str) -> Optional[Dict[str, Any]]:
    """Devuelve el registro más reciente para una IP."""
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("""
        SELECT fecha, contador, diferencia, estado
        FROM registros
        WHERE ip = ?
        ORDER BY fecha DESC
        LIMIT 1
    """, (ip,))
    row = c.fetchone()
    conn.close()
    if row:
        return {'fecha': row[0], 'contador': row[1], 'diferencia': row[2], 'estado': row[3]}
    return None

def importar_desde_excel(ruta_excel: Path) -> int:
    """
    Importa todos los registros del Excel a SQLite, calculando diferencias.
    Retorna el número de registros importados.
    """
    if not ruta_excel.exists():
        logger.error(f"El archivo {ruta_excel} no existe.")
        return 0

    try:
        wb = load_workbook(ruta_excel)
    except InvalidFileException as e:
        logger.error(f"Archivo no válido: {e}")
        return 0
    except Exception as e:
        logger.error(f"Error al cargar Excel: {e}")
        return 0

    # Buscar hoja 'Contadores'
    hoja_nombre = None
    for nombre in wb.sheetnames:
        if nombre.lower() == "contadores":
            hoja_nombre = nombre
            break
    if not hoja_nombre:
        logger.error("No se encontró la hoja 'Contadores'.")
        return 0

    ws = wb[hoja_nombre]
    logger.info(f"Hoja encontrada: {hoja_nombre}")

    col_fecha = 1
    col_ip = 4
    col_contador = 5

    # Encontrar última fila con datos
    ultima_fila = 1
    for fila in range(ws.max_row, 0, -1):
        if ws.cell(fila, col_fecha).value is not None:
            ultima_fila = fila
            break

    logger.info(f"Última fila con datos: {ultima_fila}")

    # Agrupar registros por IP
    registros_por_ip = {}
    for fila in range(2, ultima_fila + 1):
        ip = ws.cell(fila, col_ip).value
        if ip not in registros_por_ip:
            registros_por_ip[ip] = []

        fecha_celda = ws.cell(fila, col_fecha).value
        contador_celda = ws.cell(fila, col_contador).value

        # Convertir fecha
        if isinstance(fecha_celda, datetime):
            fecha_str = fecha_celda.strftime("%Y-%m-%d")
        elif isinstance(fecha_celda, str):
            try:
                fecha_dt = datetime.strptime(fecha_celda, "%d/%m/%Y")
                fecha_str = fecha_dt.strftime("%Y-%m-%d")
            except:
                continue
        else:
            continue

        # Convertir contador
        try:
            contador_int = int(contador_celda) if contador_celda is not None else None
        except:
            contador_int = None

        if contador_int is not None:
            registros_por_ip[ip].append((fecha_str, contador_int))

    # Importar cada IP
    total_importados = 0
    for ip, registros in registros_por_ip.items():
        # Ordenar por fecha
        registros.sort(key=lambda x: x[0])
        nombre = NOMBRES.get(ip, ip)
        sector = SECTORES.get(ip, "Sector no definido")

        ultimo_contador = None
        for fecha_str, contador in registros:
            if ultimo_contador is not None:
                diferencia = contador - ultimo_contador
            else:
                diferencia = None
            guardar_registro(ip, nombre, sector, fecha_str, contador, diferencia)
            ultimo_contador = contador
            total_importados += 1

    logger.info(f"Importados {total_importados} registros desde Excel a SQLite.")
    return total_importados

def obtener_resumen(fecha: str) -> Dict[str, Any]:
    """Genera un resumen de todos los registros para la fecha dada."""
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("""
        SELECT ip, nombre, sector, contador, diferencia, estado
        FROM registros
        WHERE fecha = ?
        ORDER BY ip
    """, (fecha,))
    rows = c.fetchall()
    conn.close()

    impresoras = {}
    for row in rows:
        ip, nombre, sector, contador, diferencia, estado = row
        impresoras[ip] = {
            'nombre': nombre,
            'sector': sector,
            'contador': contador,
            'diferencia': diferencia,
            'estado': estado
        }

    # Asegurar que todas las IPs configuradas estén presentes
    for ip in CONFIG.keys():
        if ip not in impresoras:
            impresoras[ip] = {
                'nombre': NOMBRES.get(ip, ip),
                'sector': SECTORES.get(ip, "Sector no definido"),
                'contador': None,
                'diferencia': None,
                'estado': 'Sin datos'
            }

    return {'fecha': fecha, 'impresoras': impresoras}