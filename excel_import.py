# excel_import.py
# Lee el archivo Excel y guarda todos los registros en SQLite

import logging
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from db_manager import init_db, guardar_registro, limpiar_base
from config import CONFIG, NOMBRES, SECTORES

logger = logging.getLogger(__name__)

def importar_excel_a_sqlite(ruta_excel: Path):
    """Importa todos los registros del Excel a SQLite, calculando diferencias."""
    if not ruta_excel.exists():
        logger.error(f"El archivo {ruta_excel} no existe.")
        return False

    try:
        wb = load_workbook(ruta_excel, data_only=True)
    except Exception as e:
        logger.error(f"Error al cargar Excel: {e}")
        return False

    # Buscar hoja 'Contadores'
    hoja = None
    for nombre in wb.sheetnames:
        if nombre.lower() == "contadores":
            hoja = wb[nombre]
            break
    if hoja is None:
        logger.error("No se encontró la hoja 'Contadores'.")
        return False

    # Columnas: A=fecha, B=sector, C=nombre, D=ip, E=contador
    col_fecha = 1
    col_ip = 4
    col_contador = 5

    # Limpiar base de datos antes de importar (opcional)
    limpiar_base()  # Borra todos los registros existentes
    init_db()

    # Recorrer filas desde la 2 hasta la última con datos
    ultima_fila = hoja.max_row
    registros_por_ip = {}

    for fila in range(2, ultima_fila + 1):
        ip = hoja.cell(fila, col_ip).value
        if ip is None or ip not in CONFIG:
            continue
        fecha_celda = hoja.cell(fila, col_fecha).value
        contador_celda = hoja.cell(fila, col_contador).value

        # Convertir fecha
        if isinstance(fecha_celda, datetime):
            fecha_str = fecha_celda.strftime("%Y-%m-%d")
        elif isinstance(fecha_celda, str):
            try:
                fecha_dt = datetime.strptime(fecha_celda, "%d/%m/%Y")
                fecha_str = fecha_dt.strftime("%Y-%m-%d")
            except:
                continue
        else:
            continue

        # Convertir contador a int o None
        try:
            contador = int(contador_celda) if contador_celda is not None else None
        except:
            contador = None

        if contador is None:
            continue

        # Agrupar por IP
        if ip not in registros_por_ip:
            registros_por_ip[ip] = []
        registros_por_ip[ip].append((fecha_str, contador))

    # Ordenar y guardar en SQLite con diferencias calculadas
    total = 0
    for ip, registros in registros_por_ip.items():
        registros.sort(key=lambda x: x[0])  # Orden ascendente por fecha
        nombre = NOMBRES.get(ip, ip)
        sector = SECTORES.get(ip, "Sector no definido")
        ultimo_contador = None
        for fecha_str, contador in registros:
            if ultimo_contador is not None:
                diferencia = contador - ultimo_contador
            else:
                diferencia = None
            guardar_registro(ip, nombre, sector, fecha_str, contador, diferencia)
            ultimo_contador = contador
            total += 1

    logger.info(f"Importación completada: {total} registros insertados en SQLite.")
    return True

if __name__ == "__main__":
    # Ejecutar solo una vez para cargar el historial
    import logging
    logging.basicConfig(level=logging.INFO)
    ruta = Path(r"C:\Users\mcescato\Desktop\Relevamiento_impresoras_vigente.xlsx")
    importar_excel_a_sqlite(ruta)